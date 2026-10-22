"""Verify base memory/documents and deferred features in a disposable installation."""

import argparse
import asyncio
import importlib.metadata
import importlib.util
import io
import os
from pathlib import Path
import subprocess
import sys
import tempfile

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from oh_my_agent_zero.cli import default_home, environment, python_in, read_state


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--models", action="store_true", help="Also download/run the actual embedding and speech models")
    args = parser.parse_args()
    home = default_home()
    print("Checking base document and memory support…", flush=True)
    state = read_state(home)
    framework = home / state["framework"]
    if os.environ.get("OMA0_FEATURE_TEST_CHILD") != "1":
        env = environment(home, state)
        env["OMA0_FEATURE_TEST_CHILD"] = "1"
        subprocess.run([str(python_in(framework)), str(Path(__file__).resolve()), *sys.argv[1:]], env=env, check=True)
        return
    from oh_my_agent_zero.native import apply
    from oh_my_agent_zero.features import ensure
    source = home / "source"
    apply(source)
    optional = ("whisper", "kokoro", "spacy", "cv2", "unstructured_inference", "triton", "llvmlite")
    base_only = not any((framework / "native-features" / (name + ".json")).exists()
                        for name in ("speech-input", "speech-output", "documents"))
    if base_only:
        assert not any(importlib.util.find_spec(name) for name in optional), "Base includes heavyweight optional libraries"
    import faiss
    import fitz
    import numpy as np
    import models
    from transformers.utils import is_timm_available, is_torchvision_available
    assert is_timm_available() and is_torchvision_available(), "Vision helpers must be discovered at initial import"
    from plugins._document_query.helpers.fetch import FetchedDocument
    from plugins._document_query.helpers.parsers import PdfParser, UnstructuredParser
    from plugins._whisper_stt.helpers import runtime as stt
    from plugins._kokoro_tts.helpers import runtime as tts

    pdf = fitz.open()
    pdf.new_page().insert_text((72, 72), "BASE_PDF_OK")
    document = FetchedDocument("native.pdf", "file", "application/pdf", pdf.tobytes())
    assert "BASE_PDF_OK" in asyncio.run(PdfParser().parse(document, {}))
    pdf.close()
    from docx import Document
    doc = Document()
    doc.add_paragraph("BASE_OFFICE_OK")
    buffer = io.BytesIO()
    doc.save(buffer)
    office_bytes = buffer.getvalue()
    document = FetchedDocument("native.docx", "file", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", office_bytes)
    assert "BASE_OFFICE_OK" in asyncio.run(UnstructuredParser().parse(document, {}, timeout=30))
    if args.models:
        embedding = models.get_embedding_model("huggingface", "sentence-transformers/all-MiniLM-L6-v2")
        vectors = np.asarray(embedding.embed_documents(["A cat sits on a chair.", "A rocket travels through space."]), dtype="float32")
    else:
        # Core imports must work before any optional feature is present.
        from sentence_transformers import SentenceTransformer
        assert callable(SentenceTransformer)
        vectors = np.eye(2, dtype="float32")
    index = faiss.IndexFlatL2(vectors.shape[1])
    index.add(vectors)
    with tempfile.TemporaryDirectory() as folder:
        path = str(Path(folder) / "memory.faiss")
        faiss.write_index(index, path)
        restored = faiss.read_index(path)
        assert restored.search(vectors[:1], 1)[1][0][0] == 0
    print("PASS: base PDF/Word extraction, local embedding interface, FAISS write/read/search.", flush=True)
    if base_only:
        assert not any(importlib.util.find_spec(name) for name in optional), "Basic documents installed the advanced stack"

    before = {d.metadata["Name"].lower(): d.version for d in importlib.metadata.distributions()}
    ensure("speech-input", source)
    import whisper
    assert callable(whisper.load_model)
    ensure("speech-output", source)
    from kokoro import KPipeline
    import spacy
    assert callable(KPipeline) and spacy.util.is_package("en_core_web_sm")

    # A real document parser triggers installation before its short parsing deadline.
    document = FetchedDocument("native.docx", "file", "application/octet-stream", office_bytes)
    assert "BASE_OFFICE_OK" in asyncio.run(UnstructuredParser().parse(document, {}, timeout=30))
    # Import availability is cached before optional installation: exercise a real table backbone.
    print("Checking table-model initialization after optional setup…", flush=True)
    import torch
    from transformers import TableTransformerConfig, TableTransformerForObjectDetection
    config = TableTransformerConfig(
        backbone="resnet18", use_pretrained_backbone=False,
        encoder_layers=1, decoder_layers=1, d_model=32,
        encoder_ffn_dim=32, decoder_ffn_dim=32, num_queries=2,
    )
    model = TableTransformerForObjectDetection(config).eval()
    print("Checking table-model inference…", flush=True)
    with torch.no_grad():
        assert model(torch.zeros(1, 3, 64, 64)).logits.shape[1] == 2
    after = {d.metadata["Name"].lower(): d.version for d in importlib.metadata.distributions()}
    assert all(after.get(name) == version for name, version in before.items()), "Optional setup changed a base dependency"
    if args.models:
        asyncio.run(stt.preload("tiny"))
        result = stt._model.transcribe(np.zeros(16000, dtype="float32"), fp16=False, language="en")
        assert isinstance(result["text"], str)
        audio = asyncio.run(tts.synthesize_sentences(["Native speech works."], {"voice": "paradee"}))
        assert audio
        print("PASS: real local embeddings, Whisper CPU inference, Paradee speech synthesis.", flush=True)
    print("PASS: all optional dependency bundles, document parsing, unchanged base versions.", flush=True)


if __name__ == "__main__":
    main()
